#!/usr/bin/env python3
"""Build a PowerPoint deck for the Threat Detection project."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_title_slide(title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle

def add_bullet_slide(title, bullets, notes=None):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    tf = slide.placeholders[1].text_frame
    for b in bullets:
        p = tf.add_paragraph()
        p.text = b
    if notes:
        slide.notes_slide.notes_text_frame.text = notes

# Title
add_title_slide("AI-Assisted Threat Detection & Response System",
                "End-to-end threat detection pipeline built on CICIDS2017\nAll offline-capable")

# Problem
add_bullet_slide("The Problem", [
    "Networks generate millions of flows daily",
    "Analysts cannot manually triage all alerts",
    "False positives waste time; missed attacks cost money",
    "Small IT teams in Lusaka, Zambia lack resources",
])

# Solution
add_bullet_slide("The Solution", [
    "ML-based detection (XGBoost) on flow data",
    "Alert correlation into incidents",
    "MITRE ATT&CK mapping for context",
    "AI SOC assistant for response recommendations",
    "Standalone demo page, fully offline",
])

# Architecture
add_bullet_slide("Architecture (6-Step Pipeline)", [
    "Step 1: Load & clean CICIDS2017 CSVs -> clean.csv",
    "Step 2: Train XGBoost with k-fold OOF -> predictions.csv",
    "Step 3: Correlate alerts into incidents -> alerts.json, incidents.json",
    "Step 4: AI assistant generates reports -> reports/*.md",
    "Step 5: Hard evaluation (chronological, LOAO, cross-capture)",
    "Step 6: Build demo HTML page",
])

# Data
add_bullet_slide("Data: CICIDS2017", [
    "957,000+ network flows",
    "80 features per flow",
    "Classes: Benign, DDoS, PortScan, Brute Force",
    "No label leakage: Destination Port excluded",
])

# Model
add_bullet_slide("ML Model", [
    "XGBoost classifier",
    "k-fold out-of-fold predictions (no train/test leakage)",
    "Baseline precision: 0.99987, recall: 0.99988",
    "FPR: 6.1e-5",
    "Comparison: with vs without Destination Port",
])

# Evaluation
add_bullet_slide("Evaluation Results", [
    "Baseline (no port): P=0.99987, R=0.99988, F1=0.99987",
    "Chronological split: P=0.9995, R=0.9970",
    "LOAO (DDoS unseen): R=0.4496",
    "LOAO (PortScan unseen): R=0.0004",
    "Honest evaluation reveals generalization gaps",
])

# MITRE
add_bullet_slide("MITRE ATT&CK Mapping", [
    "PortScan -> T1595 Active Scanning (Reconnaissance)",
    "DDoS -> T1498 Network DoS (Impact)",
    "FTP-Patator -> T1110.001 Brute Force (Credential Access)",
    "SSH-Patator -> T1110.001 Brute Force (Credential Access)",
])

# AI Assistant
add_bullet_slide("AI SOC Assistant", [
    "Offline knowledge-base mode (no API key)",
    "Optional LLM mode: Anthropic Claude, OpenRouter, Ollama",
    "Generates incident reports with response steps",
    "Tailored for small IT teams in Lusaka, Zambia",
    "Free, open-source, practical",
])

# Demo
add_bullet_slide("Demo", [
    "Single-file HTML demo: outputs/demo.html",
    "Open in any browser, no server needed",
    "Shows predictions, alerts, incidents, reports",
])

# Limitations
add_bullet_slide("Limitations", [
    "Supervised only - poor zero-shot generalization",
    "CICIDS2017 synthetic data",
    "Flow-level only - no payload/L7 visibility",
    "Single-source correlation",
    "LLM not evaluated (no API key in test env)",
])

# Future
add_bullet_slide("Future Work", [
    "Unsupervised anomaly detection layer",
    "Live capture integration (Zeek, Suricata)",
    "Graph-based multi-source correlation",
    "Edge deployment (Raspberry Pi, Jetson)",
])

# Q&A
add_title_slide("Q & A", "Thank you")

prs.save("/home/gold/GoldLevel/Threat Detection/slides.pptx")
print("Saved slides.pptx")
